# modulo_reportes/reporte_word.py
"""Generación de informe técnico detallado en formato Word (.docx) utilizando python-docx"""
import io
import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from modulo_base_datos.repositorio_arrecifes import cargar_ranking_srpi, obtener_kpis_ecosistema
import config

def generar_word_tecnico() -> bytes:
    """Genera documento Word con especificaciones de restauración arrecifal"""
    doc = Document()
    
    # Título Principal
    titulo = doc.add_heading("CoralTwin-DT — Dictamen Técnico de Restauración Arrecifal", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtítulo y Metadatos
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(f"Gemelo Digital de Arrecifes de Coral | Generado el {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()
    
    # Sección 1: Objetivos del Estudio
    h1 = doc.add_heading("1. Marco Científico y Justificación", level=1)
    p_just = doc.add_paragraph(
        "CoralTwin-DT integra modelos biofísicos de nicho ecológico acoplados con "
        "asimilación de anomalías térmicas superficiales (SST), semanas de estrés térmico (DHW de NOAA CRW) "
        "y acidificación oceánica. El objetivo es priorizar zonas donde las intervenciones de micro-fragmentación "
        "y áreas marinas protegidas maximicen la resiliencia y mantengan una cobertura coralina superior al 30% a 2050."
    )
    
    # Sección 2: Métricas de Ecosistema
    doc.add_heading("2. Diagnóstico de la Red de Monitoreo", level=1)
    kpis = obtener_kpis_ecosistema()
    
    tabla_kpis = doc.add_table(rows=1, cols=3)
    tabla_kpis.style = 'Light Shading Accent 1'
    hdr = tabla_kpis.rows[0].cells
    hdr[0].text = "Métrica Oceanográfica"
    hdr[1].text = "Valor Estimado"
    hdr[2].text = "Evaluación de Riesgo"
    
    datos = [
        ("Estaciones Totales Monitoreadas", str(kpis["total_estaciones"]), "Cobertura Global"),
        ("Cobertura Coralina Media", f"{kpis['cobertura_promedio']} %", "Vulnerable"),
        ("Semanas de Estrés Térmico (DHW)", f"{kpis['dhw_promedio']} °C-sem", "Alerta Temprana"),
        ("pH Medio Superficial", str(kpis["ph_promedio"]), "Tendencia Acidificante"),
        ("Zonas Tier 1 (Alta Prioridad)", str(kpis["prioridad_alta_restauracion"]), "Intervención Urgente")
    ]
    
    for met, val, eval_r in datos:
        row = tabla_kpis.add_row().cells
        row[0].text = met
        row[1].text = val
        row[2].text = eval_r
        
    doc.add_paragraph()
    
    # Sección 3: Priorización Espacial SRPI
    doc.add_heading("3. Priorización Espacial de Restauración (SRPI)", level=1)
    doc.add_paragraph(
        "El Índice de Priorización de Restauración Espacial (SRPI) pondera la presencia de refugios térmicos naturales, "
        "la rugosidad estructural del arrecife, la calidad del agua y la conectividad larval. Las estaciones clasificadas "
        "en 'Tier 1' representan los sitios con mayor probabilidad de supervivencia ante olas de calor marinas."
    )
    
    df_rank = cargar_ranking_srpi()
    if not df_rank.empty:
        t_srpi = doc.add_table(rows=1, cols=6)
        t_srpi.style = 'Table Grid'
        h_srpi = t_srpi.rows[0].cells
        h_srpi[0].text = "Estación"
        h_srpi[1].text = "Región"
        h_srpi[2].text = "DHW (°C-sem)"
        h_srpi[3].text = "pH"
        h_srpi[4].text = "Índice SRPI"
        h_srpi[5].text = "Zonificación"
        
        for _, r in df_rank.head(10).iterrows():
            row = t_srpi.add_row().cells
            row[0].text = str(r["Station_Name"])
            row[1].text = str(r["Region"])
            row[2].text = f"{r['DHW_degC_weeks']:.2f}"
            row[3].text = f"{r['pH_total']:.2f}"
            row[4].text = f"{r['Spatial_Restoration_Priority_Index_SRPI']:.3f}"
            row[5].text = str(r["Priority_Tier"]).replace("Tier_", "T ")
            
    doc.add_paragraph()
    
    # Pie de página y atribución
    p_pie = doc.add_paragraph(f"{config.ATTRIBUTION_TEXT} | Versión {config.VERSION}")
    p_pie.runs[0].font.size = Pt(8)
    p_pie.runs[0].font.italic = True
    
    buffer = io.BytesIO()
    doc.save(buffer)
    word_bytes = buffer.getvalue()
    buffer.close()
    return word_bytes
